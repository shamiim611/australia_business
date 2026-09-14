"""Reproduce Part 2 using published aggregates; no network or microdata required."""
import csv,json,math
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'research/part-2-ai-adoption'
RELEASE='https://www.abs.gov.au/statistics/industry/technology-and-innovation/characteristics-australian-business/2024-25'
METHOD='https://www.abs.gov.au/methodologies/characteristics-australian-business-methodology/'
SELECTION=[('2019-20','3-7',8,'ai'),('2021-22','7-1',11,'ai'),('2024-25','6-2',10,'ai'),('2024-25','6-2',0,'cloud'),('2024-25','6-2',4,'crm'),('2024-25','6-2',7,'analytics'),('2024-25','6-2',8,'erp')]
LABELS={'ai':'Artificial intelligence','cloud':'Cloud technology','crm':'Customer relationship management','analytics':'Data analytics','erp':'Enterprise resource planning'}
books={};observations=[];audit=[];verified=0
for year,tid,index,key in SELECTION:
 d=json.loads((ROOT/f'data/dashboard-{year}.json').read_text(encoding='utf-8'));t=next(t for t in d['tables'] if t['id']==tid);h=t['headers'][index]
 if t['source'] not in books:books[t['source']]=openpyxl.load_workbook(ROOT/t['source'],data_only=True)
 s=books[t['source']][t['sheet']]
 assert h['unit']=='%' and h['population']=='All businesses'
 first=min(r['row'] for r in t['rows']);lastcol=max(h['column'] for h in t['headers'])+1
 rawheaders=[{'row':r,'values':[s.cell(r,c).value for c in range(1,lastcol+1)]} for r in range(1,first)]
 gate='C' if year!='2019-20' else 'D'
 reason='Cross-sectional use approved; no temporal comparison approved. Survey framework or response context changed.' if gate=='C' else 'Historical snapshot only: AI examples include voice recognition; full item equivalence not established.'
 record=f'{year}:{tid}:{index}'
 audit.append({'record_id':record,'year':year,'variable':key,'original_wording':h['label'],'examples':json.dumps(h.get('notes',[]),ensure_ascii=False),'response_options':json.dumps([x['label'] for x in t['headers']],ensure_ascii=False),'denominator':h['denominator'],'unit':'%','business_scope':'Employing businesses with ABS exclusions; agriculture included.','size_categories':json.dumps([r['label'] for r in t['rows'] if r['group']=='Employment size'],ensure_ascii=False),'industry_classification':'ANZSIC 2006 (Revision 2.0)','reference_period':t.get('period',d.get('period',year)),'comparable_years':'None approved for temporal analysis','temporal_gate':gate,'within_wave':'Usable with published reliability flags','rationale':reason,'workbook':t['source'],'sheet':t['sheet'],'raw_headers':json.dumps(rawheaders,ensure_ascii=False),'source_notes':json.dumps(t['notes'],ensure_ascii=False),'methodology':METHOD+year,'reviewed':'2026-09-14'})
 for r in t['rows']:
  if r['label']=='0-19 persons':continue
  c=r['cells'][index];raw=s.cell(c['sourceRow'],c['sourceColumn']+1);value=c['value']
  if isinstance(value,(int,float)) and isinstance(raw.value,(int,float)):
   assert math.isclose(value,raw.value,abs_tol=1e-9),(record,r['label'],value,raw.value)
   verified+=1
  elif value is not None:assert c['status']=='rounded-zero' and value==0
  observations.append({'year':year,'variable':key,'label':LABELS[key],'group':r['group'],'category':r['label'],'value':value,'unit':'%','status':c['status'],'note':c['note'],'denominator':h['denominator'],'record_id':record,'workbook':t['source'],'sheet':t['sheet'],'cell':c['cell'],'source_number_format':raw.number_format,'source_value':raw.value,'temporal_gate':gate})
def valid(o,strict=False):return o['value'] is not None and o['status'] in (['published','rounded-zero'] if strict else ['published','rounded-zero','caution'])
ai=[o for o in observations if o['year']=='2024-25' and o['variable']=='ai']
national=next(o for o in ai if o['group']=='Overall')
size=[o for o in ai if o['group']=='Employment size'];assert len(size)==4
large=size[3];gaps=[]
for small,label in [(size[0],'Large minus micro'),(size[1],'Large minus small')]:
 gaps.append({'comparison':label,'year':'2024-25','large_category':large['category'],'comparison_category':small['category'],'large_percent':large['value'],'comparison_percent':small['value'],'gap_pp':round(large['value']-small['value'],1) if all(valid(o) for o in [large,small]) else None,'unflagged_inputs':all(valid(o,True) for o in [large,small]),'large_cell':large['cell'],'comparison_cell':small['cell']})
industries=[o for o in ai if o['group']=='Industry'];dispersion=[]
for strict in [False,True]:
 used=[o for o in industries if valid(o,strict)];lo=min(used,key=lambda o:o['value']);hi=max(used,key=lambda o:o['value'])
 dispersion.append({'sample':'Exclude all flagged inputs' if strict else 'Include usable cautions','industries':len(used),'minimum_industry':lo['category'],'minimum_percent':lo['value'],'maximum_industry':hi['category'],'maximum_percent':hi['value'],'range_pp':round(hi['value']-lo['value'],1)})
benchmarks=[{'industry':o['category'],'ai_percent':o['value'],'national_percent':national['value'],'gap_pp':round(o['value']-national['value'],1) if valid(o) and valid(national) else None,'status':o['status'],'cell':o['cell']} for o in industries]
evidence={'innovation_status':'D - withheld from derived comparisons','source':RELEASE,'reviewed':'2026-09-14','finding':'The release narrative describes higher AI adoption among innovation-active businesses. Its chart footnote instead describes a denominator of businesses reporting AI use. The published chart and narrative do not resolve this inconsistency.','decision':'No innovation gap or causal inference is calculated. No numeric innovation chart is included until the denominator is confirmed.','history':'The release names 2022-23 in its AI key statistics and 2021-22 in the detailed ICT discussion. Local workbook periods identify the historical snapshots used here. No continuous trend or growth calculation is approved.'}
result={'observations':observations,'audit':audit,'gaps':gaps,'dispersion':dispersion,'industry_benchmarks':benchmarks,'national':national,'verified_cells':verified,'evidence':evidence,'labels':LABELS}
OUT.mkdir(parents=True,exist_ok=True)
def savecsv(name,rows):
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
for name,rows in [('observations.csv',observations),('available_ai_observations.csv',ai),('comparability_decisions.csv',audit),('size_gaps.csv',gaps),('industry_benchmarks.csv',benchmarks),('industry_dispersion.csv',dispersion)]:savecsv(name,rows)
(OUT/'evidence_decisions.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
(OUT/'results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'data/research-ai.js').write_text('window.AI_RESEARCH = '+json.dumps(result,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
f=['# Part 2: AI adoption is unevenly distributed','','Status: completed descriptive public-data analysis. Temporal and innovation-status comparisons remain unapproved where source evidence is insufficient.','','## Findings',f"In 2024-25, {national['value']}% of businesses in the survey scope reported AI use. The item is broad AI use, not specifically generative AI or intensity of use. Source: {national['workbook']}, {national['sheet']}, {national['cell']}.",'']
f += [f"- {o['category']}: {o['value']}% ({o['cell']})." for o in size]
f += ['',*[f"{g['comparison']}: {g['gap_pp']} percentage points; excluding flagged inputs leaves the comparison unchanged: {g['unflagged_inputs']}." for g in gaps],'',f"The {len(industries)} industry rates range from {dispersion[0]['minimum_percent']}% in {dispersion[0]['minimum_industry']} to {dispersion[0]['maximum_percent']}% in {dispersion[0]['maximum_industry']}: {dispersion[0]['range_pp']} points. The unflagged-only range is {dispersion[1]['range_pp']} points. These are differences in group adoption rates, not the distribution of AI adopters across industries.",'','The Northern Territory estimate carries a caution. Geographic differences are unadjusted for business composition; overlapping location categories are not additive.','', 'Cloud, CRM, analytics and ERP provide same-wave context. Separate marginal shares do not measure co-adoption or the capabilities of AI users.','','## Historical and innovation evidence',evidence['history'],'',evidence['finding'],evidence['decision'],f"See the [ABS release]({RELEASE}) for the external evidence reviewed.",'','## Limits','This is a cross-section of aggregate estimates, not a firm-level panel. No causal claims, significance tests, regression, synthetic industry-by-size cells, weighted SME rates or maturity index are produced. Exact uncertainty for these AI comparisons has not been established; source flags are retained. Rankings do not prove statistically distinct performance.','','## Reproduce',f"Run `python scripts/build_ai_research.py` then `python scripts/test_ai_research.py`. The builder verifies {verified} numeric workbook cells and writes source-linked observations, item decisions, differences and sensitivity results. No network access is needed to rebuild. Open `research-ai.html` offline."]
(OUT/'findings.md').write_text('\n'.join(f)+'\n',encoding='utf-8')
print(f'Built Part 2: {len(observations)} observations; {verified} source numbers verified; {len(audit)} audited items.')
